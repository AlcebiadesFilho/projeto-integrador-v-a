"""
Gera os slides da apresentação (.pptx) com a mesma identidade visual do painel.

Uso: python scripts/build_slides.py
Pré-requisito: figuras em docs/figuras (python scripts/capture_figures.py)
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIGS = ROOT / "docs" / "figuras"
OUT = ROOT / "entregaveis" / "Apresentacao_PI5A_Infomach.pptx"

FONT = "Helvetica Neue"
PAGE = RGBColor(0xF4, 0xF3, 0xEF)
SURFACE = RGBColor(0xFC, 0xFC, 0xFB)
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x89, 0x87, 0x81)
HAIR = RGBColor(0xE1, 0xE0, 0xD9)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
ORANGE_WASH = RGBColor(0xF6, 0xE6, 0xDD)
YELLOW = RGBColor(0xFF, 0xF2, 0x9E)

W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.75)  # margem lateral


# ---------------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------------
def bg(slide, color=PAGE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def text(slide, x, y, w, h, runs, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1):
    """runs: str | list[str | (str, dict)] — dict aceita color/bold/size. '\n' quebra parágrafo."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [runs]
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    for r in runs:
        t, o = (r, {}) if isinstance(r, str) else r
        parts = t.split("\n")
        for k, part in enumerate(parts):
            if k:
                p = tf.add_paragraph()
                p.alignment = align
                p.line_spacing = spacing
            if not part:
                continue
            run = p.add_run()
            run.text = part
            f = run.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.color.rgb = o.get("color", color)
            if o.get("highlight"):
                rPr = run._r.get_or_add_rPr()
                hl = rPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}highlight", {})
                clr = hl.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr", {"val": "FFF29E"})
                hl.append(clr)
                rPr.append(hl)
    return tb


def box(slide, x, y, w, h, fill=SURFACE, line=HAIR, radius=0.06):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    return s


def rect(slide, x, y, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def picture(slide, name, x, y, max_w, max_h):
    path = FIGS / f"{name}.png"
    with Image.open(path) as im:
        ratio = im.height / im.width
    w = max_w
    h = int(w * ratio)
    if h > max_h:
        h = max_h
        w = int(h / ratio)
    # centraliza no espaço disponível
    px = x + (max_w - w) // 2
    return slide.shapes.add_picture(str(path), px, y, Emu(w), Emu(h))


def header(slide, n, kicker, title, page):
    bg(slide)
    text(slide, MX, Inches(0.5), Inches(6), Inches(0.3), [(f"{n}  ", {"color": ORANGE, "bold": True}), kicker.upper()],
         size=11, color=INK2)
    text(slide, MX, Inches(0.85), Inches(11.8), Inches(1.2), title, size=32, bold=True, spacing=1.0)
    text(slide, W - MX - Inches(3), H - Inches(0.45), Inches(3), Inches(0.25), f"PI V-A · Infomach   {page}",
         size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def finding(slide, x, y, w, h, runs, size=16, label="ACHADO"):
    box(slide, x, y, w, h, fill=ORANGE_WASH, line=None, radius=0.08)
    rect(slide, x, y + Inches(0.12), Inches(0.06), h - Inches(0.24), ORANGE)
    text(slide, x + Inches(0.3), y + Inches(0.18), w - Inches(0.5), Inches(0.25), label, size=10, bold=True, color=INK2)
    text(slide, x + Inches(0.3), y + Inches(0.45), w - Inches(0.5), h - Inches(0.55), runs, size=size, spacing=1.15)


def stat(slide, x, y, w, value, label, accent=False, vsize=40):
    text(slide, x, y, w, Inches(0.8), value, size=vsize, bold=True, color=ORANGE if accent else INK)
    text(slide, x, y + Inches(0.75), w, Inches(0.7), label, size=13, color=INK2)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


# ---------------------------------------------------------------------------------
# slides
# ---------------------------------------------------------------------------------
def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    page = [0]

    def new():
        page[0] += 1
        return prs.slides.add_slide(blank)

    # 1 — capa
    s = new()
    bg(s)
    rect(s, 0, 0, Inches(0.18), H, ORANGE)
    text(s, MX, Inches(0.7), Inches(10), Inches(0.3),
         "PUC GOIÁS · CEAD · BIG DATA E INTELIGÊNCIA ARTIFICIAL · PROJETO INTEGRADOR V-A", size=11, color=INK2)
    text(s, MX, Inches(2.0), Inches(11), Inches(2.6),
         ["O gargalo do service desk\nnão é esforço. ", ("É fila.", {"color": ORANGE})], size=60, bold=True, spacing=0.95)
    text(s, MX, Inches(4.75), Inches(10), Inches(0.9),
         "Painel de inteligência operacional sobre a base do Movidesk da Infomach:\nindicadores, visualização e data storytelling",
         size=18, color=INK2, spacing=1.2)
    text(s, MX, Inches(6.3), Inches(8), Inches(0.6),
         [("Alcebiades Alves Barbosa Filho", {"bold": True, "color": INK}), "\nGoiânia, outubro de 2026"], size=13, color=INK2)
    notes(s, "Abertura: apresentar o projeto e a tese em uma frase.")

    # 2 — organização parceira
    s = new()
    header(s, "01", "Organização parceira", "Infomach: segurança e TI gerenciadas para centenas de empresas", page[0])
    text(s, MX, Inches(2.2), Inches(5.6), Inches(3.5),
         "A Infomach é um provedor de serviços gerenciados de TI e Segurança (MSSP) de Goiânia/GO.\n\n"
         "Todo atendimento técnico, de incidentes e alertas a configurações e reuniões, é registrado no Movidesk, "
         "e cada técnico aponta as horas trabalhadas em cada ação.\n\n"
         "Esta é a 3ª etapa da parceria: Data Lake (PI IV-A) → temas com NLP (PI IV-B) → BI e visualização (PI V-A).",
         size=16, color=INK2, spacing=1.25)
    teams = ["Rede & Conectividade", "Backup", "Endpoint & Proteção de Dados", "Cloud", "Infraestrutura", "Monitoramento"]
    box(s, Inches(7.0), Inches(2.2), Inches(5.6), Inches(4.2))
    text(s, Inches(7.35), Inches(2.45), Inches(5), Inches(0.3), "EQUIPES TÉCNICAS", size=11, bold=True, color=INK2)
    for i, t in enumerate(teams):
        y = Inches(2.95 + i * 0.55)
        rect(s, Inches(7.35), y + Inches(0.1), Inches(0.12), Inches(0.12), ORANGE if t == "Backup" else BLUE)
        text(s, Inches(7.65), y, Inches(4.8), Inches(0.4), t, size=16)
    notes(s, "Contexto da organização e do setor envolvido (operações / service desk).")

    # 3 — problema
    s = new()
    header(s, "02", "Necessidade identificada", "O Movidesk registra cada chamado. Ninguém via o conjunto.", page[0])
    text(s, MX, Inches(2.15), Inches(11.5), Inches(0.8),
         "Os relatórios nativos são listas operacionais: mostram chamados, não mostram tendência, concentração nem gargalo. "
         "Com a coordenação técnica, a necessidade virou quatro perguntas:", size=16, color=INK2, spacing=1.2)
    qs = [("A demanda cresce? O esforço acompanha?", "Dimensionar equipes"),
          ("Quanto é reativo e onde ele trava?", "Onde atacar primeiro"),
          ("A prioridade muda o tempo de atendimento?", "Régua de prioridade / SLA"),
          ("Para onde vão as horas técnicas?", "Contratos e alocação")]
    for i, (q, d) in enumerate(qs):
        x = MX + i * Inches(3.0)
        box(s, x, Inches(3.4), Inches(2.8), Inches(2.9))
        text(s, x + Inches(0.3), Inches(3.65), Inches(2.3), Inches(0.4), f"0{i + 1}", size=14, bold=True, color=ORANGE)
        text(s, x + Inches(0.3), Inches(4.1), Inches(2.3), Inches(1.4), q, size=18, bold=True, spacing=1.1)
        text(s, x + Inches(0.3), Inches(5.6), Inches(2.3), Inches(0.5), [("Decisão: ", {"color": MUTED}), d], size=12, color=INK2)
    text(s, MX, Inches(6.55), Inches(11.5), Inches(0.4),
         "Público: diretoria e coordenação técnica, líderes de equipe, Comercial / Customer Success.", size=13, color=MUTED)
    notes(s, "Necessidade, público e decisões que o painel apoia.")

    # 4 — interação
    s = new()
    header(s, "03", "Interação com a organização", "Construído com a Infomach, não só sobre os dados dela", page[0])
    box(s, MX, Inches(2.15), Inches(5.6), Inches(2.9))
    text(s, MX + Inches(0.35), Inches(2.4), Inches(5), Inches(0.3), "INTERLOCUTOR", size=11, bold=True, color=INK2)
    text(s, MX + Inches(0.35), Inches(2.85), Inches(5), Inches(0.6), "Murillo Santos", size=30, bold=True)
    text(s, MX + Inches(0.35), Inches(3.55), Inches(5), Inches(0.4), "COO · Diretor de Operações da Infomach", size=16, color=INK2)
    text(s, MX + Inches(0.35), Inches(4.25), Inches(5), Inches(0.6), "Público principal do painel: a diretoria de operações.",
         size=13, color=MUTED)
    steps = [("Etapas anteriores", "Cessão da base anonimizada do Movidesk (Termo de Cessão, PI IV-A / IV-B)", BLUE),
             ("30/09/2026", "Reunião com o COO: painel, KPIs, achados e recomendações", ORANGE)]
    for i, (d, t, c) in enumerate(steps):
        y = Inches(2.15) + i * Inches(1.5)
        x = MX + Inches(6.1)
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, x, y + Inches(0.08), Inches(0.26), Inches(0.26))
        dot.fill.solid()
        dot.fill.fore_color.rgb = c
        dot.line.fill.background()
        if i == 0:
            rect(s, x + Inches(0.11), y + Inches(0.4), Inches(0.04), Inches(1.15), HAIR)
        text(s, x + Inches(0.5), y, Inches(5.2), Inches(0.4), d, size=17, bold=True)
        text(s, x + Inches(0.5), y + Inches(0.42), Inches(5.2), Inches(0.9), t, size=14, color=INK2, spacing=1.15)
    finding(s, MX, Inches(5.4), Inches(11.8), Inches(1.4),
            "Retorno muito positivo: valorizou a visão consolidada que o Movidesk não oferece, apontou a fila do Backup e a "
            "régua de prioridade como ações imediatas e quer usar o painel nas reuniões da coordenação.", size=15,
            label="FEEDBACK DO COO")
    notes(s, "Contar como foi a reunião de 30/09 com o COO Murillo Santos e o retorno positivo.")

    # 5 — dados
    s = new()
    header(s, "04", "Base de dados", "95 mil apontamentos de horas viraram 41 mil chamados analisáveis", page[0])
    for i, (v, l) in enumerate([("95.291", "apontamentos de horas\n(23 colunas, Movidesk)"), ("41.676", "chamados na janela\njul/2020 – mai/2022"),
                                ("51.281 h", "de trabalho técnico\napontado"), ("326", "clientes · 52 técnicos\n7 equipes")]):
        stat(s, MX + i * Inches(3.0), Inches(2.3), Inches(2.8), v, l, vsize=38)
    box(s, MX, Inches(4.2), Inches(11.8), Inches(2.5))
    text(s, MX + Inches(0.35), Inches(4.4), Inches(5.5), Inches(0.3), "TRATAMENTOS (PYTHON / PANDAS)", size=11, bold=True, color=INK2)
    text(s, MX + Inches(0.35), Inches(4.8), Inches(5.6), Inches(1.8),
         "• Equipes unificadas (grafias duplicadas, renomeação)\n• Categorias corrigidas e tipo de demanda criado\n"
         "• Tempo de ciclo: abertura → último apontamento", size=14, color=INK2, spacing=1.3)
    text(s, MX + Inches(6.2), Inches(4.4), Inches(5.3), Inches(0.3), "PRIVACIDADE E QUALIDADE", size=11, bold=True, color=INK2)
    text(s, MX + Inches(6.2), Inches(4.8), Inches(5.3), Inches(1.8),
         "• Sem texto livre, nomes ou e-mails (LGPD)\n• Clientes e técnicos anonimizados na origem\n"
         "• Meses parciais excluídos das tendências", size=14, color=INK2, spacing=1.3)
    notes(s, "Origem, atributos e tratamentos. Lembrar que o ciclo é um proxy, pois não há data de fechamento.")

    # 6 — KPIs
    s = new()
    header(s, "05", "Indicadores (KPIs)", "Nove indicadores, cada um ligado a uma decisão", page[0])
    kpis = [("1.812", "chamados / mês", "dimensionamento"), ("2.230 h", "horas técnicas / mês", "custo operacional"),
            ("74 min", "esforço médio / chamado", "custo por chamado"), ("94 h", "horas / técnico / mês", "capacidade"),
            ("44,6%", "demanda reativa", "prevenção / automação"), ("15,5 h", "ciclo mediano", "tempo percebido"),
            ("30,7%", "urgente + alta", "régua de prioridade"), ("21,3%", "horas em atividade interna", "capacidade não faturada"),
            ("31,0%", "horas nos 10 maiores clientes", "revisão de contratos")]
    for i, (v, l, d) in enumerate(kpis):
        r, c = divmod(i, 3)
        x, y = MX + c * Inches(3.97), Inches(2.1) + r * Inches(1.6)
        box(s, x, y, Inches(3.8), Inches(1.45))
        accent = i in (4, 5)
        text(s, x + Inches(0.3), y + Inches(0.18), Inches(3.3), Inches(0.6), v, size=28, bold=True, color=ORANGE if accent else INK)
        text(s, x + Inches(0.3), y + Inches(0.78), Inches(3.3), Inches(0.3), l, size=13, color=INK)
        text(s, x + Inches(0.3), y + Inches(1.05), Inches(3.3), Inches(0.3), f"→ {d}", size=11, color=MUTED)
    notes(s, "Cada KPI responde a uma decisão. Todos recalculam com os filtros do painel.")

    # 7 — solução
    s = new()
    header(s, "06", "A solução", "Um painel web que conta a história, do problema à ação", page[0])
    picture(s, "painel_hero", MX, Inches(2.05), Inches(7.2), Inches(4.7))
    x = MX + Inches(7.6)
    text(s, x, Inches(2.1), Inches(4.2), Inches(0.3), "ARQUITETURA", size=11, bold=True, color=INK2)
    flow = [("Movidesk", "CSV exportado"), ("Python · pandas", "limpeza e KPIs"), ("JSON anonimizado", "1,3 MB"),
            ("React + Recharts", "filtros no navegador")]
    for i, (a, b) in enumerate(flow):
        y = Inches(2.55) + i * Inches(0.95)
        box(s, x, y, Inches(4.2), Inches(0.75), fill=SURFACE)
        text(s, x + Inches(0.25), y + Inches(0.12), Inches(3.8), Inches(0.5), [(a, {"bold": True}), f"  ·  {b}"], size=14, color=INK2)
        if i < len(flow) - 1:
            text(s, x + Inches(1.95), y + Inches(0.72), Inches(0.4), Inches(0.25), "↓", size=12, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, x, Inches(6.4), Inches(4.3), Inches(0.6), "Open source · site estático · modo claro/escuro · responsivo",
         size=12, color=MUTED)
    notes(s, "Por que front-end próprio: controle do design (tema é visualização), narrativa em página única, publicação gratuita, filtros instantâneos. Python continua como motor dos dados.")

    # 8 — escolhas de visualização
    s = new()
    header(s, "07", "Escolhas de visualização", "Cada gráfico escolhido pela pergunta que responde", page[0])
    rules = [("Cor com função", "Azul = a série. Laranja só para o problema. Cinza = contexto. Paleta validada para daltonismo."),
             ("Um eixo por gráfico", "Chamados × horas indexados (base 100), nunca eixo duplo."),
             ("Título-conclusão", "O título diz a mensagem; o gráfico prova."),
             ("Menos ruído", "Grade fina, sem 3D, rótulos só onde ajudam."),
             ("Acessível", "Legenda sempre, tabela “Dados” em cada gráfico, contraste e movimento reduzido.")]
    for i, (t, d) in enumerate(rules):
        y = Inches(2.1) + i * Inches(0.95)
        text(s, MX, y, Inches(5.2), Inches(0.35), t, size=16, bold=True)
        text(s, MX, y + Inches(0.37), Inches(5.2), Inches(0.55), d, size=12.5, color=INK2, spacing=1.1)
    picture(s, "reativo_por_equipe", MX + Inches(5.8), Inches(2.0), Inches(6.0), Inches(4.8))
    notes(s, "Mostrar ênfase: Backup em laranja, demais em cinza. Explicar a validação da paleta.")

    # 9 — narrativa
    s = new()
    header(s, "08", "Data storytelling", "Uma página, um arco: contexto → conflito → resolução", page[0])
    arc = [("Tese", "“O gargalo não é esforço. É fila.”"), ("00 · Problema", "4 perguntas, 4 decisões"),
           ("01 · Visão geral", "KPIs e tendência"), ("02 · Reativo", "Conflito: fila no Backup"),
           ("03 · Prioridade", "Conflito: Alta = Normal"), ("04 · Esforço", "Conflito: interno e Pareto"),
           ("05 · Ações", "Resolução + KPI de cada ação")]
    for i, (t, d) in enumerate(arc):
        x = MX + i * Inches(1.7)
        conflict = 3 <= i <= 5
        box(s, x, Inches(2.4), Inches(1.55), Inches(2.4), fill=ORANGE_WASH if conflict else SURFACE, line=None if conflict else HAIR)
        text(s, x + Inches(0.15), Inches(2.6), Inches(1.3), Inches(0.8), t, size=14, bold=True, spacing=1.05)
        text(s, x + Inches(0.15), Inches(3.5), Inches(1.3), Inches(1.2), d, size=11.5, color=INK2, spacing=1.1)
    finding(s, MX, Inches(5.3), Inches(11.8), Inches(1.45),
            "Os textos de “Achado” são calculados a partir dos dados: ao filtrar o painel, os números e a própria conclusão "
            "mudam. Se o recorte contradiz a história geral, o texto diz isso.", size=15)
    notes(s, "Explicar a ordem dos capítulos e a barra de filtros fixa.")

    # 10 — insight 1
    s = new()
    header(s, "09", "Insight 1 · Demanda", "Chamados cresceram 66%. As horas, só 19%.", page[0])
    picture(s, "indice_demanda_horas", MX, Inches(2.0), Inches(8.2), Inches(4.9))
    x = MX + Inches(8.6)
    stat(s, x, Inches(2.2), Inches(3.2), "+66%", "chamados por mês\n(3 primeiros × 3 últimos meses)", accent=True)
    stat(s, x, Inches(4.0), Inches(3.2), "+19%", "horas apontadas\nno mesmo período")
    text(s, x, Inches(5.7), Inches(3.2), Inches(1.0), "Os chamados novos são mais curtos: alertas e incidentes de rotina.",
         size=13, color=INK2, spacing=1.2)
    notes(s, "Índice base 100 para evitar eixo duplo.")

    # 11 — insight 2
    s = new()
    header(s, "10", "Insight 2 · Demanda reativa", "Reativo: 45% dos chamados, 15% das horas, 3,7× mais tempo na fila", page[0])
    picture(s, "mix_demanda", MX, Inches(2.0), Inches(11.8), Inches(1.6))
    picture(s, "ciclo_backup_trimestral", MX, Inches(3.85), Inches(7.6), Inches(3.1))
    x = MX + Inches(8.0)
    stat(s, x, Inches(3.95), Inches(3.8), "48%", "de todo o trabalho reativo é do Backup", accent=True, vsize=44)
    text(s, x, Inches(5.35), Inches(3.8), Inches(1.4),
         "13 min de esforço por chamado, 28,5 h de ciclo mediano. 76% vêm do CloudBackup.", size=14, color=INK2, spacing=1.2)
    notes(s, "O problema é fila, não esforço técnico. O Backup oscila em ondas.")

    # 12 — insight 3
    s = new()
    header(s, "11", "Insight 3 · Prioridade", "Marcar como “Alta” não acelera o atendimento", page[0])
    picture(s, "ciclo_por_prioridade", MX, Inches(2.0), Inches(7.4), Inches(4.9))
    x = MX + Inches(7.9)
    for i, (v, l, acc) in enumerate([("5 h", "Urgente", False), ("22 h", "Alta", True), ("23 h", "Normal", False)]):
        stat(s, x, Inches(2.15) + i * Inches(1.25), Inches(3.8), v, f"ciclo mediano · {l}", accent=acc, vsize=34)
    text(s, x, Inches(6.0), Inches(3.9), Inches(0.9),
         "Ou “Alta” é atribuída sem critério, ou não há fila priorizada para ela.", size=13, color=INK2, spacing=1.2)
    notes(s, "Só Urgente se diferencia. Recomendação: régua objetiva + fila própria.")

    # 13 — insight 4
    s = new()
    header(s, "12", "Insight 4 · Esforço e clientes", "Um quinto das horas fica em casa. 27 clientes levam metade do resto.", page[0])
    picture(s, "pareto_clientes", MX, Inches(2.0), Inches(7.4), Inches(4.9))
    x = MX + Inches(7.9)
    stat(s, x, Inches(2.15), Inches(3.8), "21%", "das horas em atividade interna\n(7.405 h só em reuniões)", accent=True)
    stat(s, x, Inches(3.9), Inches(3.8), "8%", "dos clientes externos consomem\n50% das horas (27 de 325)")
    text(s, x, Inches(5.7), Inches(3.9), Inches(1.0), "Cliente_4 tratado como conta interna (inferido; a confirmar com a Infomach).",
         size=11.5, color=MUTED, spacing=1.2)
    notes(s, "Pareto clássico. Base para revisão de contratos e técnico dedicado.")

    # 14 — recomendações
    s = new()
    header(s, "13", "Recomendações", "O que os dados sugerem fazer, e como medir", page[0])
    recs = [("Automatizar a triagem dos alertas de Backup", "ciclo mediano reativo do Backup"),
            ("Revisar a régua de prioridade (critério + fila para “Alta”)", "ciclo por prioridade"),
            ("Medir e orçar a atividade interna", "% de horas internas"),
            ("Revisar contratos do topo do Pareto", "% de horas nos 10 maiores"),
            ("Registrar primeira resposta e fechamento no Movidesk", "% de chamados com SLA medido")]
    for i, (t, k) in enumerate(recs):
        y = Inches(2.05) + i * Inches(0.95)
        box(s, MX, y, Inches(11.8), Inches(0.8))
        text(s, MX + Inches(0.3), y + Inches(0.2), Inches(0.6), Inches(0.4), f"0{i + 1}", size=16, bold=True, color=ORANGE)
        text(s, MX + Inches(1.0), y + Inches(0.2), Inches(7.0), Inches(0.4), t, size=16, bold=True)
        text(s, MX + Inches(8.2), y + Inches(0.23), Inches(3.4), Inches(0.4), [("medir: ", {"color": MUTED}), k], size=13, color=INK2)
    notes(s, "Cada recomendação tem o KPI que mostra se funcionou.")

    # 15 — contribuições e extensão
    s = new()
    header(s, "14", "Contribuições e experiência extensionista", "O que a Infomach leva, e o que eu levo", page[0])
    box(s, MX, Inches(2.1), Inches(5.7), Inches(4.6))
    text(s, MX + Inches(0.35), Inches(2.35), Inches(5), Inches(0.3), "PARA A INFOMACH", size=11, bold=True, color=INK2)
    text(s, MX + Inches(0.35), Inches(2.8), Inches(5.1), Inches(3.8),
         "• Visão de gestão que não existia: tendência, gargalo e concentração\n• Painel atualizável com dois comandos\n"
         "• 5 ações priorizadas, cada uma com um KPI\n• Lacunas de registro mapeadas (SLA, apontamentos retroativos)",
         size=15, color=INK2, spacing=1.35)
    box(s, MX + Inches(6.1), Inches(2.1), Inches(5.7), Inches(4.6), fill=ORANGE_WASH, line=None)
    text(s, MX + Inches(6.45), Inches(2.35), Inches(5), Inches(0.3), "EXPERIÊNCIA EXTENSIONISTA", size=11, bold=True, color=INK2)
    text(s, MX + Inches(6.45), Inches(2.8), Inches(5.1), Inches(3.8),
         "• Traduzir perguntas de gestão em indicadores que apoiam decisões reais\n"
         "• Mostrar menos e dizer mais: um gráfico, uma mensagem\n"
         "• Cuidar do dado de terceiros: anonimização e LGPD desde o início\n"
         "• Apresentar para a diretoria, não para técnicos, e ouvir o retorno",
         size=15, color=INK2, spacing=1.35)
    notes(s, "Fechar com o valor para a organização e o aprendizado pessoal.")

    # 16 — encerramento
    s = new()
    bg(s)
    rect(s, 0, 0, Inches(0.18), H, ORANGE)
    text(s, MX, Inches(2.3), Inches(11), Inches(1.4), "Obrigado.", size=60, bold=True)
    text(s, MX, Inches(3.8), Inches(11), Inches(1.5),
         [("Código e painel: ", {"color": MUTED}), "github.com/AlcebiadesFilho/projeto-integrador-v-a",
          ("\nPara rodar: ", {"color": MUTED}), "cd web  →  npm install  →  npm run dev"], size=18, spacing=1.4)
    text(s, MX, Inches(6.3), Inches(10), Inches(0.6),
         [("Alcebiades Alves Barbosa Filho", {"bold": True, "color": INK}), "  ·  PUC Goiás · Big Data e IA · 2026"], size=13, color=INK2)

    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)
    print("[ok]", OUT.relative_to(ROOT), f"({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
